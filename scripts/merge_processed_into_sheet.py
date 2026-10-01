#!/usr/bin/env python3
"""
Carry the processed-image work from an older sheet onto a newer one.

The founder edits the master sheet while images are being processed, so the
new download has rows added or removed and nothing lines up by position. This
copies, per product:

    Resized Aniket      the Drive folder of converted JPEGs
    SKU Code            the SKU already baked into those folder names
    Variant Option      Colour / State, where the row has variants
    Variant Values      SLUG=Label pairs

Rows are matched on the "Product Images - Drive folder link" cell (column AD),
the list of source filenames. On the real sheets that key matched 239 of 239
rows with no duplicates on either side, where S. No. matched 238 and titles
collided 27 times.

Nothing is re-uploaded and no SKU is re-issued: the Drive folders are named
after the old SKUs, so they are copied across rather than recomputed.

Install
-------
    pip install openpyxl

Usage
-----
    python3 merge_processed_into_sheet.py NEW.xlsx OLD-processed.xlsx
    python3 merge_processed_into_sheet.py NEW.xlsx OLD-processed.xlsx --out ready.xlsx
    python3 merge_processed_into_sheet.py NEW.xlsx OLD-processed.xlsx --dry-run
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

SHEET_NAME = "Main Product Sheet | Aniket"
HEADER_ROW = 2
FIRST_DATA_ROW = 3

SOURCE_IMAGES_HEADER = "Product Images - Drive folder link"
BRIEF_HEADER = "Brief Product Description"
TITLE_HEADER = "Product Title - only if different from Brief Description"
SKU_HEADER = "SKU Code"

# What the processed link column is called. The first name is what this script
# writes; the uploader accepts either.
TARGET_LINK_HEADER = "Resized Aniket"
LEGACY_LINK_HEADERS = ["Processed Images - Drive folder link"]

VARIANT_OPTION_HEADER = "Variant Option"
VARIANT_VALUES_HEADER = "Variant Values"
DEV_NOTES_HEADER = "Developer Notes"

STATE_WORDS = {
    "mizoram", "manipur", "tripura", "meghalaya", "nagaland", "sikkim",
    "assam", "arunachal", "pradesh", "bengal", "bihar", "odisha", "assamese",
}

# An option name is a word or two: "Colour", "State", "Design", "Fabric".
# Anything longer is someone's note, not an option.
PLAUSIBLE_OPTION = re.compile(r"^[A-Za-z][A-Za-z /-]{0,19}$")


def looks_like_option(text: str) -> bool:
    t = str(text or "").strip()
    return bool(t) and len(t.split()) <= 2 and bool(PLAUSIBLE_OPTION.match(t))


def derive_option(values_cell: str) -> str:
    """Works the option name out from the values, ignoring the stored one."""
    labels = []
    for pair in str(values_cell or "").split("|"):
        i = pair.find("=")
        if i != -1:
            labels.append(pair[i + 1:].strip())
    if not labels:
        return "Colour"
    hits = sum(1 for l in labels
               for w in re.findall(r"[A-Za-z]+", l)
               if w.lower() in STATE_WORDS)
    return "State" if hits >= len(labels) else "Colour"


def column_indexes(ws):
    cols = {}
    for c in range(1, ws.max_column + 2):
        v = ws.cell(HEADER_ROW, c).value
        if v and str(v).strip():
            cols[str(v).strip()] = c
    return cols


def open_tab(path: Path, tab: str):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    if tab not in wb.sheetnames:
        sys.exit("Tab %r not found in %s. Tabs: %s" % (tab, path.name, wb.sheetnames))
    return wb, wb[tab]


def cell_text(ws, row, col):
    if not col:
        return ""
    v = ws.cell(row, col).value
    return "" if v is None else str(v).strip()


def norm_key(text: str) -> str:
    """
    Normalises an image-list cell so trivial edits still match: case, repeated
    whitespace, and trailing separators are ignored.
    """
    t = re.sub(r"\s+", " ", str(text or "")).strip().lower()
    return t.strip(" ,;")


def product_rows(ws, cols):
    brief_col = cols.get(BRIEF_HEADER)
    out = []
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        brief = cell_text(ws, row, brief_col)
        if not brief or brief == "Filled":
            continue
        out.append(row)
    return out


def label_for(ws, row, cols):
    t = cell_text(ws, row, cols.get(TITLE_HEADER))
    return t or cell_text(ws, row, cols.get(BRIEF_HEADER))


def main() -> int:
    ap = argparse.ArgumentParser(
        description="Copy processed-image links, SKUs and variant columns from "
                    "an older processed sheet onto a freshly downloaded one.")
    ap.add_argument("new_sheet", help="the newly downloaded .xlsx")
    ap.add_argument("processed_sheet", help="the older -processed.xlsx")
    ap.add_argument("--out", help="output path (default: <new>-ready.xlsx)")
    ap.add_argument("--tab", default=SHEET_NAME)
    ap.add_argument("--dry-run", action="store_true",
                    help="report the match and write nothing")
    args = ap.parse_args()

    new_path = Path(args.new_sheet)
    old_path = Path(args.processed_sheet)
    for p in (new_path, old_path):
        if not p.exists():
            sys.exit("No such file: %s" % p)

    new_wb, new_ws = open_tab(new_path, args.tab)
    _old_wb, old_ws = open_tab(old_path, args.tab)

    new_cols = column_indexes(new_ws)
    old_cols = column_indexes(old_ws)

    src_col_new = new_cols.get(SOURCE_IMAGES_HEADER)
    src_col_old = old_cols.get(SOURCE_IMAGES_HEADER)
    if not src_col_new or not src_col_old:
        sys.exit("Both sheets need a %r column." % SOURCE_IMAGES_HEADER)

    # Where the old sheet keeps the processed link — either name.
    old_link_col = None
    for header in [TARGET_LINK_HEADER] + LEGACY_LINK_HEADERS:
        if old_cols.get(header):
            old_link_col = old_cols[header]
            break
    if not old_link_col:
        sys.exit("%s has no processed-link column (looked for %s)."
                 % (old_path.name,
                    ", ".join([TARGET_LINK_HEADER] + LEGACY_LINK_HEADERS)))

    # ── index the old sheet ──────────────────────────────────
    old_by_key, old_dupes = {}, set()
    for row in product_rows(old_ws, old_cols):
        key = norm_key(cell_text(old_ws, row, src_col_old))
        if not key:
            continue
        if key in old_by_key:
            old_dupes.add(key)
            continue
        old_by_key[key] = {
            "row": row,
            "link": cell_text(old_ws, row, old_link_col),
            "sku": cell_text(old_ws, row, old_cols.get(SKU_HEADER)),
            "vopt": cell_text(old_ws, row, old_cols.get(VARIANT_OPTION_HEADER)),
            "vval": cell_text(old_ws, row, old_cols.get(VARIANT_VALUES_HEADER)),
            "label": label_for(old_ws, row, old_cols),
        }

    def ensure_column(header):
        if header in new_cols:
            return new_cols[header]
        c = max(new_cols.values()) + 1
        new_ws.cell(HEADER_ROW, c, header)
        new_cols[header] = c
        print("Added column %d: %s" % (c, header))
        return c

    col_notes = new_cols.get(DEV_NOTES_HEADER)
    if not args.dry_run:
        col_link = ensure_column(TARGET_LINK_HEADER)
        col_sku = ensure_column(SKU_HEADER)
        col_vopt = ensure_column(VARIANT_OPTION_HEADER)
        col_vval = ensure_column(VARIANT_VALUES_HEADER)
        col_notes = new_cols.get(DEV_NOTES_HEADER)
        print("")

    matched, no_link, unmatched, variant_rows = 0, [], [], []
    rescued_notes, fixed_options = [], []
    used = set()

    for row in product_rows(new_ws, new_cols):
        key = norm_key(cell_text(new_ws, row, src_col_new))
        hit = old_by_key.get(key) if key else None

        if not hit:
            unmatched.append((row, label_for(new_ws, row, new_cols)))
            continue

        used.add(key)
        if not hit["link"]:
            no_link.append((row, label_for(new_ws, row, new_cols)))
            continue

        matched += 1

        # Someone used the Variant Option column for photo notes ("fingers are
        # showing"). Never pass that through as a product option name: derive
        # the name from the values, and keep the note somewhere useful.
        stored = hit["vopt"]
        option = ""
        note = ""
        if hit["vval"]:
            if looks_like_option(stored):
                option = stored
            else:
                option = derive_option(hit["vval"])
                if stored:
                    note = stored
                    fixed_options.append((row, hit["sku"], stored, option))
        elif stored and not looks_like_option(stored):
            note = stored

        if note:
            rescued_notes.append((row, hit["sku"], note))

        if hit["vval"]:
            variant_rows.append((row, hit["sku"], option,
                                 hit["vval"].count("|") + 1))

        if not args.dry_run:
            new_ws.cell(row, col_link, hit["link"])
            if hit["sku"]:
                new_ws.cell(row, col_sku, hit["sku"])
            if option:
                new_ws.cell(row, col_vopt, option)
            if hit["vval"]:
                new_ws.cell(row, col_vval, hit["vval"])
            if note and col_notes:
                existing = cell_text(new_ws, row, col_notes)
                if note not in existing:
                    new_ws.cell(row, col_notes,
                                (existing + " | " + note) if existing else note)

    dropped = [(v["row"], v["label"]) for k, v in old_by_key.items()
               if k not in used]

    out = Path(args.out) if args.out else new_path.with_name(
        new_path.stem + "-ready.xlsx")
    if not args.dry_run:
        try:
            new_wb.save(out)
        except PermissionError:
            sys.exit("Could not write %s — is it open in Excel?" % out.name)

    # ── report ──────────────────────────────────────────────
    print("=" * 64)
    print("Matched and filled     : %d" % matched)
    print("Variant products       : %d" % len(variant_rows))
    print("Matched but no link    : %d" % len(no_link))
    print("In new sheet, no match : %d" % len(unmatched))
    print("Removed by the founder : %d" % len(dropped))
    if old_dupes:
        print("Duplicate keys in old  : %d (first occurrence used)" % len(old_dupes))
    if not args.dry_run:
        print("Written                : %s" % out)

    if fixed_options:
        print("\nVariant Option held a note, not an option name. Derived the")
        print("option from the values instead:")
        for row, sku, was, now in fixed_options:
            print("   row %-4d %-14s %r -> %s" % (row, sku, was[:40], now))

    if rescued_notes:
        print("\nPhoto notes found in the Variant Option column, copied to")
        print("%s so they are not lost:" % DEV_NOTES_HEADER)
        for row, sku, note in rescued_notes:
            print("   row %-4d %-14s %s" % (row, sku, note[:56]))

    if variant_rows:
        print("\nVariant products carried over:")
        for row, sku, vopt, n in variant_rows:
            print("   row %-4d %-16s %-7s %d values" % (row, sku, vopt, n))

    if dropped:
        print("\nProducts in the processed sheet that are gone from the new one")
        print("(their Drive folders stay put, nothing is deleted):")
        for row, label in dropped:
            print("   was row %-4d %s" % (row, label[:60]))

    if no_link:
        print("\nMatched but the old sheet had no processed link — run")
        print("prepare_product_images.py for these:")
        for row, label in no_link:
            print("   row %-4d %s" % (row, label[:60]))

    if unmatched:
        print("\nNo match in the processed sheet. These are new products, or")
        print("their image list was edited — run prepare_product_images.py:")
        for row, label in unmatched:
            print("   row %-4d %s" % (row, label[:60]))

    return 0


if __name__ == "__main__":
    raise SystemExit(main())

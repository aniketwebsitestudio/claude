#!/usr/bin/env python3
"""
One pass over the product sheet that gets every image Shopify-ready.

For each product row it:
  1. reads the image filenames (column AD) and the Drive folder link (AE)
  2. downloads those files from Drive (cached locally, so re-runs are cheap)
  3. converts RAW (.CR3/.NEF/...) to JPEG and resizes anything over Shopify's
     25 MP / 20 MB caps
  4. renames each file to the product's SKU plus a sequence number
  5. uploads the results to a fresh per-product folder in Drive
  6. writes the SKU and the new folder link back into a copy of the sheet

Upload that copy to Google Sheets and the Apps Script uploader takes it from
there — it pulls every image straight out of the processed folder, so nothing
has to be matched by filename again.

SKUs follow the Bhoomija scheme: [COLLECTION]-[SUBCATEGORY]-[SEQUENCE],
e.g. HOM-PLN-0001.

Install
-------
    pip install rawpy pillow openpyxl \
        google-api-python-client google-auth-httplib2 google-auth-oauthlib

Credentials
-----------
Put credentials.json (OAuth Desktop client) next to this script, same as your
other Drive scripts. token.json is created on first run.

Usage
-----
    python3 prepare_product_images.py sheet.xlsx --create-parent "Bhoomija Shopify Images"
    python3 prepare_product_images.py sheet.xlsx --parent <DRIVE_FOLDER_ID>
    python3 prepare_product_images.py sheet.xlsx --parent <ID> --rows 4-20
    python3 prepare_product_images.py sheet.xlsx --dry-run

Drive layout it produces:

    <parent>/
      ACC-JHL-0001_Naga Jhola Bag/
        ACC-JHL-0001_01.jpg
        ACC-JHL-0001_02.jpg
      WEA-DUP-0001_Teal green handwoven cotton.../
        WEA-DUP-0001_01.jpg
        ...
"""

from __future__ import annotations

import argparse
import io
import os
import re
import sys
from pathlib import Path

BASE_DIR = Path(__file__).parent
CREDENTIALS_FILE = BASE_DIR / "credentials.json"
TOKEN_FILE = BASE_DIR / "token.json"
CACHE_DIR = BASE_DIR / "_image_cache"
SCOPES = ["https://www.googleapis.com/auth/drive"]

SHEET_NAME = "Main Product Sheet | Aniket"
HEADER_ROW = 2
FIRST_DATA_ROW = 3

LONG_EDGE = 2048
QUALITY = 88
MAX_PIXELS = 25_000_000
MAX_BYTES = 20 * 1024 * 1024

# Matches MAX_IMAGES_PER_PRODUCT in the Apps Script. Uploading more than the
# uploader will read just wastes bandwidth and Drive quota.
MAX_IMAGES = 10

NEW_COLUMN_HEADER = "Processed Images - Drive folder link"

RAW_EXT = {".cr3", ".cr2", ".crw", ".nef", ".nrw", ".arw", ".srf", ".sr2",
           ".dng", ".raf", ".orf", ".rw2", ".pef", ".kdc", ".dcr", ".x3f"}
NATIVE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".gif", ".tif", ".tiff"}

# Sub-category spellings in the sheet that differ from the canonical name.
# Kept in step with SUBCATEGORY_ALIAS in bhoomija-shopify-uploader.gs.
SUB_ALIAS = {
    ("Home", "Durrie"): "Durrie / Throw",
    ("Home", "Fabric Throw"): "Durrie / Throw",
    ("Wearables", "Mekhela Chador"): "Mekhla Chador",
    ("Wearables", "Wrapper"): "Mekhla Wrapper",
    ("Wall Art Forms", "Wall Hanger"): "Embroidered Wall Hanger",
}

# The SKU scheme splits the store's single "Gift Ideas" collection into
# Bags & Carry and Home & Art, so the collection code depends on the
# sub-category for that one category.
GIFT_BAGS = {"Bag", "Slingbag", "Jhola", "Clutch", "Pouch", "Hamper",
             "Jewellery Box"}

COLLECTION_CODE = {
    "Home": "HOM",
    "Wearables": "WEA",
    "Jewellery": "JWL",
    "Wall Art Forms": "WAR",
    "Accessories": "ACC",
    "Special Feature": "SPF",   # not in the scheme doc; added here
}

SUB_CODE = {
    "Figurine": "FIG", "Lamp": "LMP", "Planter": "PLN", "Cushion Cover": "CSC",
    "Bed Spread": "BSP", "Coaster": "CST", "Floor Mat": "FLM", "Mat Set": "MTS",
    "Placemat": "PMT", "Durrie / Throw": "DUR", "Mirror": "MIR",
    "Runner": "RUN", "Tissue Holder": "TSH",
    "Dupatta": "DUP", "Stole": "STL", "Mekhla Wrapper": "MKW",
    "Mekhla Chador": "MKC", "Fabric": "FAB",
    "Necklace": "NCK", "Earrings": "ERG", "Hairpin": "HRP", "Bala": "BAL",
    "Bag": "BAG", "Slingbag": "SLB", "Jhola": "JHL", "Clutch": "CLT",
    "Pouch": "POU", "Hamper": "HMP", "Jewellery Box": "JBX",
    "Tray": "TRY", "Wall Hanger": "WHG", "Serviette": "SRV",
    "Table top frame": "TTF", "Stand": "STD", "Holder": "HLD",
    "Painting": "PNT", "Sculpture": "SCU", "Bookmark": "BKM", "Magnet": "MAG",
    "Embroidered Wall Hanger": "EWH",
    "Belt": "BLT", "Thaila": "THL",
    "Decor": "DEC",   # not in the scheme doc; added here
}


# ── SKU ──────────────────────────────────────────────────────────────────

def canonical_sub(category: str, sub: str) -> str:
    return SUB_ALIAS.get((category, sub), sub)


def codes_for(category: str, sub: str):
    """Returns (collection_code, sub_code) or (None, reason)."""
    sub = canonical_sub(category, sub)

    if category == "Gift Ideas":
        collection = "GBC" if sub in GIFT_BAGS else "GHA"
    else:
        collection = COLLECTION_CODE.get(category)

    if not collection:
        return None, "unknown category %r" % category
    code = SUB_CODE.get(sub)
    if not code:
        return None, "unknown sub-category %r under %r" % (sub, category)
    return collection, code


class SkuAllocator:
    """4-digit running number per collection+subcategory, in sheet order."""

    def __init__(self):
        self.counts: dict[tuple[str, str], int] = {}

    def next(self, collection: str, sub_code: str) -> str:
        key = (collection, sub_code)
        self.counts[key] = self.counts.get(key, 0) + 1
        return "%s-%s-%04d" % (collection, sub_code, self.counts[key])


# ── image work ───────────────────────────────────────────────────────────

def fit(size, long_edge):
    w, h = size
    longest = max(w, h)
    if longest <= long_edge:
        return None
    scale = long_edge / float(longest)
    return max(1, int(round(w * scale))), max(1, int(round(h * scale)))


def load_image(path: Path):
    """Open any supported file as a PIL RGB image. Returns (img, how)."""
    from PIL import Image

    if path.suffix.lower() in RAW_EXT:
        import rawpy
        with rawpy.imread(str(path)) as raw:
            full_long = max(raw.sizes.width, raw.sizes.height)
            try:
                thumb = raw.extract_thumb()
                if thumb.format == rawpy.ThumbFormat.JPEG:
                    img = Image.open(io.BytesIO(thumb.data))
                    img.load()
                    if max(img.size) >= min(full_long, LONG_EDGE):
                        return img.convert("RGB"), "embedded preview"
                elif thumb.format == rawpy.ThumbFormat.BITMAP:
                    img = Image.fromarray(thumb.data)
                    if max(img.size) >= min(full_long, LONG_EDGE):
                        return img.convert("RGB"), "embedded preview"
            except Exception:
                pass
            rgb = raw.postprocess(use_camera_wb=True, no_auto_bright=False,
                                  output_bps=8)
            return Image.fromarray(rgb), "developed from raw"

    img = Image.open(path)
    img.load()
    # Honour the camera's rotation flag before we throw the EXIF away.
    try:
        from PIL import ImageOps
        img = ImageOps.exif_transpose(img)
    except Exception:
        pass
    return img.convert("RGB"), "as shot"


def to_shopify_jpeg(src: Path, dest: Path) -> str:
    """Convert/resize src into a Shopify-safe JPEG at dest. Returns a note."""
    from PIL import Image

    img, how = load_image(src)
    original = img.size

    target = fit(img.size, LONG_EDGE)
    if target:
        img = img.resize(target, Image.LANCZOS)
    while img.size[0] * img.size[1] > MAX_PIXELS:
        img = img.resize((img.size[0] // 2, img.size[1] // 2), Image.LANCZOS)

    q = QUALITY
    while True:
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=q, optimize=True, progressive=True,
                 subsampling="4:2:0")
        if buf.tell() <= MAX_BYTES or q <= 60:
            break
        q -= 8

    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(buf.getvalue())

    return "%dx%d->%dx%d %.1fMB %s" % (
        original[0], original[1], img.size[0], img.size[1],
        buf.tell() / 1048576.0, how)


# ── Drive ────────────────────────────────────────────────────────────────

def _import_google():
    try:
        from googleapiclient.discovery import build
        from googleapiclient.http import MediaFileUpload, MediaIoBaseDownload
        from google_auth_oauthlib.flow import InstalledAppFlow
        from google.auth.transport.requests import Request
        from google.oauth2.credentials import Credentials
        return build, MediaFileUpload, MediaIoBaseDownload, InstalledAppFlow, Request, Credentials
    except ImportError as exc:
        sys.exit("Missing Google libraries. Install with:\n"
                 "  pip install google-api-python-client google-auth-httplib2 "
                 "google-auth-oauthlib")


def get_drive_service():
    build, _M, _D, InstalledAppFlow, Request, Credentials = _import_google()
    creds = None
    if TOKEN_FILE.exists():
        try:
            creds = Credentials.from_authorized_user_file(str(TOKEN_FILE), SCOPES)
        except Exception:
            creds = None   # unreadable or from a different client

    if creds and not creds.valid and creds.expired and creds.refresh_token:
        try:
            creds.refresh(Request())
        except Exception as exc:
            # The saved refresh token is dead — revoked, expired, or issued by
            # an OAuth client that no longer exists. Log in again rather than
            # dying, which is what "invalid_grant: Bad Request" means.
            print("Saved login is no longer valid (%s). Signing in again..."
                  % type(exc).__name__)
            try:
                TOKEN_FILE.unlink()
            except OSError:
                pass
            creds = None

    if not creds or not creds.valid:
        if not CREDENTIALS_FILE.exists():
            sys.exit("ERROR: %s not found next to the script."
                     % CREDENTIALS_FILE.name)
        flow = InstalledAppFlow.from_client_secrets_file(
            str(CREDENTIALS_FILE), SCOPES)
        creds = flow.run_local_server(port=0)

    TOKEN_FILE.write_text(creds.to_json(), encoding="utf-8")
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def cell_link(ws, row: int, col: int):
    """
    The Drive link lives as a cell hyperlink, not as text — the visible value
    is a label like "Accessories Sept 10". Falls back to a plain URL or a
    HYPERLINK() formula.
    """
    cell = ws.cell(row, col)
    if cell.hyperlink is not None and cell.hyperlink.target:
        return cell.hyperlink.target
    value = str(cell.value or "")
    m = re.search(r'HYPERLINK\(\s*"([^"]+)"', value, re.I)
    if m:
        return m.group(1)
    return value if value.startswith("http") else ""


def extract_folder_id(url: str):
    if not url:
        return None
    for pattern in (r"/folders/([A-Za-z0-9_-]+)",
                    r"/file/d/([A-Za-z0-9_-]+)",
                    r"[?&]id=([A-Za-z0-9_-]+)"):
        m = re.search(pattern, str(url))
        if m:
            return m.group(1)
    return None


def list_folder(service, folder_id: str) -> dict[str, dict]:
    """{lowercase filename: {id, name, size}} for one folder."""
    out, page = {}, None
    while True:
        res = service.files().list(
            q="'%s' in parents and trashed = false" % folder_id,
            fields="nextPageToken, files(id, name, size, mimeType)",
            pageSize=1000, pageToken=page,
            supportsAllDrives=True, includeItemsFromAllDrives=True,
        ).execute()
        for f in res.get("files", []):
            out[f["name"].lower()] = f
        page = res.get("nextPageToken")
        if not page:
            break
    return out


def download(service, file_id: str, name: str) -> Path:
    _b, _M, MediaIoBaseDownload, *_ = _import_google()
    CACHE_DIR.mkdir(exist_ok=True)
    dest = CACHE_DIR / ("%s_%s" % (file_id[:12], name))
    if dest.exists() and dest.stat().st_size:
        return dest

    req = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, req, chunksize=8 * 1024 * 1024)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    dest.write_bytes(buf.getvalue())
    return dest


def find_or_create_folder(service, name: str, parent_id: str) -> str:
    safe = name.replace("\\", "\\\\").replace("'", "\\'")
    res = service.files().list(
        q=("name = '%s' and mimeType = 'application/vnd.google-apps.folder' "
           "and '%s' in parents and trashed = false" % (safe, parent_id)),
        fields="files(id)", pageSize=10,
        supportsAllDrives=True, includeItemsFromAllDrives=True,
    ).execute()
    files = res.get("files", [])
    if files:
        return files[0]["id"]
    created = service.files().create(
        body={"name": name,
              "mimeType": "application/vnd.google-apps.folder",
              "parents": [parent_id]},
        fields="id", supportsAllDrives=True,
    ).execute()
    return created["id"]


def upload_jpeg(service, local: Path, parent_id: str) -> None:
    _b, MediaFileUpload, *_ = _import_google()
    media = MediaFileUpload(str(local), mimetype="image/jpeg", resumable=True)
    service.files().create(
        body={"name": local.name, "parents": [parent_id]},
        media_body=media, fields="id", supportsAllDrives=True,
    ).execute()


# ── sheet ────────────────────────────────────────────────────────────────

_SANITIZE = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def sanitize(text: str, max_len: int = 60) -> str:
    text = _SANITIZE.sub("", str(text)).strip().strip(".")
    text = re.sub(r"\s+", " ", text)
    return text[:max_len].rstrip() or "Untitled"


def parse_image_names(raw: str) -> tuple[list[str], list[str]]:
    """Splits column AD into filenames, keeping RAW names."""
    if not raw:
        return [], []
    known = "|".join(e.lstrip(".") for e in sorted(NATIVE_EXT | RAW_EXT))
    pattern = re.compile(r"([^\s,][^,]*?\.(?:%s))\b" % known, re.I)
    files, unusable, seen = [], [], set()
    for entry in re.split(r"[,\n;]+", str(raw)):
        entry = entry.strip()
        if not entry:
            continue
        m = pattern.search(entry)
        if not m:
            unusable.append(entry)
            continue
        name = m.group(1).strip()
        # A few rows list the same file twice; Shopify would show it twice too.
        if name.lower() in seen:
            continue
        seen.add(name.lower())
        files.append(name)
    return files, unusable


def column_indexes(ws):
    """{header text: 1-based column} from the header row."""
    cols = {}
    for c in range(1, ws.max_column + 2):
        v = ws.cell(HEADER_ROW, c).value
        if v and str(v).strip():
            cols[str(v).strip()] = c
    return cols


def parse_rows_arg(spec: str, lo: int, hi: int) -> set[int]:
    wanted = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            wanted.update(range(int(a), int(b) + 1))
        else:
            wanted.add(int(part))
    return {r for r in wanted if lo <= r <= hi}


# ── main ─────────────────────────────────────────────────────────────────

def main() -> int:
    global LONG_EDGE

    ap = argparse.ArgumentParser(
        description="Resize, convert and re-upload product images, and write "
                    "SKUs plus processed-folder links back into the sheet.")
    ap.add_argument("sheet", help="the .xlsx product sheet")
    ap.add_argument("--parent", help="ID of an existing Drive folder to create "
                                     "the per-product folders inside")
    ap.add_argument("--create-parent", metavar="NAME",
                    help="make the parent folder instead of passing an ID. "
                         "Created at your Drive root, or inside --parent when "
                         "both are given. Re-used if a folder of that name is "
                         "already there.")
    ap.add_argument("--out", help="output .xlsx (default: <sheet>-processed.xlsx)")
    ap.add_argument("--rows", help="only these sheet rows, e.g. 4-20 or 4,7,9")
    ap.add_argument("--tab", default=SHEET_NAME, help="worksheet name")
    ap.add_argument("--long-edge", type=int, default=LONG_EDGE)
    ap.add_argument("--max-images", type=int, default=MAX_IMAGES,
                    help="images per product (default %d, matching the Apps "
                         "Script)" % MAX_IMAGES)
    ap.add_argument("--dry-run", action="store_true",
                    help="assign SKUs and report the plan; touch nothing")
    ap.add_argument("--no-upload", action="store_true",
                    help="convert locally into ./_processed but skip Drive")
    args = ap.parse_args()

    LONG_EDGE = args.long_edge

    import openpyxl

    src = Path(args.sheet)
    if not src.exists():
        sys.exit("No such sheet: %s" % src)
    if not (args.dry_run or args.no_upload or args.parent or args.create_parent):
        sys.exit("Pass --parent <FOLDER_ID> or --create-parent <NAME> "
                 "(or --dry-run / --no-upload).")

    wb = openpyxl.load_workbook(src)
    if args.tab not in wb.sheetnames:
        sys.exit("Tab %r not found. Tabs: %s" % (args.tab, wb.sheetnames))
    ws = wb[args.tab]

    cols = column_indexes(ws)
    for required in ("Categories", "Sub Category", "Brief Product Description",
                     "Product Images - Drive folder link", "SKU Code"):
        if required not in cols:
            sys.exit("Column %r not found in row %d." % (required, HEADER_ROW))

    col_cat = cols["Categories"]
    col_sub = cols["Sub Category"]
    col_brief = cols["Brief Product Description"]
    col_title = cols.get("Product Title - only if different from Brief Description")
    col_names = cols["Product Images - Drive folder link"]
    col_link = col_names + 1          # column AE, as in the Apps Script
    col_sku = cols["SKU Code"]

    col_new = cols.get(NEW_COLUMN_HEADER)
    if not col_new:
        col_new = max(cols.values()) + 1
        ws.cell(HEADER_ROW, col_new, NEW_COLUMN_HEADER)
        print("Added column %d: %s\n" % (col_new, NEW_COLUMN_HEADER))

    last = ws.max_row
    wanted = parse_rows_arg(args.rows, FIRST_DATA_ROW, last) if args.rows else None

    service = None
    parent_id = args.parent
    if not (args.dry_run or args.no_upload):
        print("Authenticating with Google Drive...")
        service = get_drive_service()

        if args.create_parent:
            parent_id = find_or_create_folder(
                service, args.create_parent, args.parent or "root")
            print("Parent folder: %s\n  https://drive.google.com/drive/folders/%s\n"
                  % (args.create_parent, parent_id))

    out = Path(args.out) if args.out else src.with_name(src.stem + "-processed.xlsx")

    def save_progress():
        """Write the sheet now, so a crash or a dropped connection cannot
        lose links for products that are already uploaded."""
        if args.dry_run:
            return
        try:
            wb.save(out)
        except PermissionError:
            print("    (could not save %s — is it open in Excel?)" % out.name)

    alloc = SkuAllocator()
    folder_index: dict[str, dict] = {}

    done = skipped = failed_images = 0
    problems: list[str] = []

    for row in range(FIRST_DATA_ROW, last + 1):
        brief = ws.cell(row, col_brief).value
        if not brief or str(brief).strip() in ("", "Filled"):
            continue

        category = str(ws.cell(row, col_cat).value or "").strip()
        sub = str(ws.cell(row, col_sub).value or "").strip()
        if not category:
            continue

        collection, code = codes_for(category, sub)
        if collection is None:
            problems.append("Row %d: %s" % (row, code))
            continue

        # Allocate for every valid row so numbering stays stable regardless
        # of which rows this run happens to process.
        sku = alloc.next(collection, code)

        if wanted is not None and row not in wanted:
            continue

        title = str(ws.cell(row, col_title).value or "").strip() if col_title else ""
        title = title or str(brief).strip()

        ws.cell(row, col_sku, sku)

        names, unusable = parse_image_names(ws.cell(row, col_names).value)
        folder_id = extract_folder_id(cell_link(ws, row, col_link))

        if len(names) > args.max_images:
            problems.append("Row %d (%s): %d images listed, keeping the first %d"
                            % (row, sku, len(names), args.max_images))
            names = names[:args.max_images]

        # A failed sheet formula leaves its source text in the cell, which
        # would otherwise become the product title in Shopify.
        if "__xludf" in title or title.startswith("="):
            problems.append("Row %d (%s): title is a broken formula, fix the "
                            "sheet" % (row, sku))

        print("[row %d] %s  %s" % (row, sku, title[:50]))

        if not names:
            skipped += 1
            problems.append("Row %d (%s): no image filenames" % (row, sku))
            print("    - no image filenames")
            continue
        if not folder_id:
            skipped += 1
            problems.append("Row %d (%s): no Drive folder link" % (row, sku))
            print("    - no Drive folder link")
            continue

        if args.dry_run:
            print("    %d image(s): %s" % (len(names), ", ".join(names)))
            done += 1
            continue

        out_dir = Path("_processed") / ("%s_%s" % (sku, sanitize(title, 40)))
        out_dir.mkdir(parents=True, exist_ok=True)

        if folder_id not in folder_index:
            try:
                folder_index[folder_id] = list_folder(service, folder_id) \
                    if service else {}
            except Exception as exc:
                problems.append("Row %d (%s): folder unreadable: %s"
                                % (row, sku, exc))
                print("    X folder unreadable: %s" % exc)
                skipped += 1
                continue
        index = folder_index[folder_id]

        produced: list[Path] = []
        for i, name in enumerate(names, start=1):
            key = name.lower()
            entry = index.get(key)
            if not entry:
                stem = os.path.splitext(key)[0]
                # A JPEG sibling beats the RAW original.
                for cand, meta in index.items():
                    if os.path.splitext(cand)[0] != stem:
                        continue
                    if os.path.splitext(cand)[1] in NATIVE_EXT:
                        entry = meta
                        break
                    entry = entry or meta
            if not entry:
                failed_images += 1
                problems.append("Row %d (%s): %s not in Drive" % (row, sku, name))
                print("    X %s  not found in Drive" % name)
                continue

            dest = out_dir / ("%s_%02d.jpg" % (sku, i))
            if dest.exists():
                produced.append(dest)
                print("    - %s  already converted" % dest.name)
                continue

            try:
                local = download(service, entry["id"], entry["name"])
                note = to_shopify_jpeg(local, dest)
                produced.append(dest)
                print("    . %s  <- %s  (%s)" % (dest.name, entry["name"], note))
            except Exception as exc:
                failed_images += 1
                problems.append("Row %d (%s): %s failed: %s: %s"
                                % (row, sku, name, type(exc).__name__, exc))
                print("    X %s  %s: %s" % (name, type(exc).__name__, exc))

        if not produced:
            skipped += 1
            continue

        if args.no_upload:
            done += 1
            continue

        target_name = "%s_%s" % (sku, sanitize(title, 40))
        try:
            drive_folder = find_or_create_folder(service, target_name, parent_id)
            already = list_folder(service, drive_folder)
            for f in produced:
                if f.name.lower() in already:
                    continue
                upload_jpeg(service, f, drive_folder)
        except Exception as exc:
            # Usually the connection dropping. The converted files are on disk
            # and whatever reached Drive stays there, so a re-run picks this
            # row up where it stopped.
            skipped += 1
            problems.append("Row %d (%s): upload failed: %s: %s"
                            % (row, sku, type(exc).__name__, exc))
            print("    X upload failed: %s: %s" % (type(exc).__name__, exc))
            save_progress()
            continue

        url = "https://drive.google.com/drive/folders/%s" % drive_folder
        ws.cell(row, col_new, url)
        print("    -> %s" % url)
        done += 1
        save_progress()

    save_progress()

    print("\n" + "=" * 62)
    print("Rows processed   : %d" % done)
    print("Rows skipped     : %d" % skipped)
    print("Images failed    : %d" % failed_images)
    if not args.dry_run:
        print("Sheet written    : %s" % out)
    if problems:
        report = Path("_processed") / "problems.txt"
        report.parent.mkdir(exist_ok=True)
        report.write_text("\n".join(problems), encoding="utf-8")
        print("Problems (%d)     : %s" % (len(problems), report))
        for p in problems[:15]:
            print("   " + p)
        if len(problems) > 15:
            print("   ... and %d more" % (len(problems) - 15))

    if not args.dry_run and not args.no_upload and done:
        print("\nNext: upload %s to Google Sheets, then run\n"
              "Bhoomija -> Setup -> Use this sheet, and upload as usual." % out.name)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
